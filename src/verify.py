# -*- coding: utf-8 -*-
# 数字自检：报告里写死的数，和 data/processed 里的实际产物对一遍。
#
# 改了数据忘了改报告是最容易出的错（这一轮就出过两次），所以把这个检查写成脚本，
# 改完任何一步都跑一下：
#
#   python src/verify.py
#
# 检查方式是"报告 docx 里有没有这个字符串"，不是去解析表格，
# 所以要求报告正文确实把关键数写出来了；没写的会提示缺失。

import json
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PROC = os.path.join(ROOT, "data", "processed")
OUT = os.path.join(ROOT, "outputs")
NAME = "第X组_组长姓名_《数据挖掘》探索性实验报告_项目1"
# 正式文件可能被预览面板或 Word/WPS 占着写不进去，脚本会另存成"（新）""（新2）"，
# 凡是同名前缀的 docx 都算候选，取修改时间最新的那个来查
_cands = [os.path.join(OUT, f) for f in os.listdir(OUT)
          if f.startswith(NAME) and f.endswith(".docx")] if os.path.isdir(OUT) else []
_cands = [c for c in _cands if os.path.exists(c)]
REPORT = max(_cands, key=os.path.getmtime) if _cands else os.path.join(OUT, NAME + ".docx")


def load(name):
    with open(os.path.join(PROC, name), encoding="utf-8") as f:
        return json.load(f)


def read_report():
    try:
        import docx
    except ImportError:
        return None
    if not os.path.exists(REPORT):
        return None
    d = docx.Document(REPORT)
    # 表格里的数字也要算进去，很多关键数是写在表里的
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def main():
    docs = load("documents.json")
    e1 = load("entities_v1.json")
    e2 = load("entities_v2.json")
    routes = load("routes.json")
    coref = load("coref_result.json")["stats"]
    conds = load("conditions.json")
    graph = load("graph.json")
    runs = json.load(open(os.path.join(OUT, "agent_runs.json"), encoding="utf-8"))

    def cnt(data, t):
        n = 0
        for r in data:
            for e in r["entities"]:
                if e["type"] == t:
                    n += 1
        return n

    checks = [
        ("文档篇数", str(len(docs))),
        ("清洗后总字数", str(sum(d["clean_char_len"] for d in docs))),
        ("v1 景点出现次数", str(cnt(e1, "POI"))),
        ("v2 景点出现次数", str(cnt(e2, "POI"))),
        ("v2 交通次数", str(cnt(e2, "TRANSPORT"))),
        ("v2 时间次数", str(cnt(e2, "TIME"))),
        ("路线条数", str(len(routes))),
        ("路线总步数", str(sum(r["step_count"] for r in routes))),
        ("条件建议条数", str(len(conds))),
        ("共指标注例数", str(coref["total"])),
        ("共指可判定例数", str(coref["clear"])),
        ("图谱节点数", str(len(graph["nodes"]))),
        ("图谱边数", str(len(graph["links"]))),
        ("测试场景数", str(len(runs))),
    ]

    text = read_report()
    bad = 0
    for label, val in checks:
        if text is None:
            print("  ? %-12s %s（报告不存在，只列出实际值）" % (label, val))
            continue
        hit = val in text
        if not hit:
            bad += 1
        print("  %s %-12s %s" % ("√" if hit else "×", label, val))

    if text is not None:
        print()
        if bad:
            print("有 %d 项在报告里找不到，改数据后记得同步正文。" % bad)
        else:
            print("报告里的关键数和产物一致。")
    return 0 if bad == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
